"""Execute in the container only; recalculate and validate XLSX before export."""
import pathlib
import subprocess
import sys
import tempfile

code = sys.stdin.read()
subprocess.run([sys.executable, "-I", "-c", code], check=True, timeout=24)
from openpyxl import load_workbook

for file in pathlib.Path('/workspace/outputs').glob('*.xlsx'):
    if file.is_symlink():
        raise RuntimeError('Symlink output rejected')
    book = load_workbook(file, read_only=True)
    if book._external_links:
        raise RuntimeError('External workbook links require manual review')
    has_formulas = any(cell.data_type == 'f' for sheet in book for row in sheet for cell in row)
    book.close()
    if has_formulas:
        with tempfile.TemporaryDirectory() as directory:
            subprocess.run(['soffice', '-env:UserInstallation=file:///tmp/lo-profile', '--headless', '--convert-to', 'xlsx', '--outdir', directory, str(file)], check=True, timeout=20, stdout=sys.stderr)
            recalculated = pathlib.Path(directory) / file.name
            if not recalculated.is_file():
                raise RuntimeError('LibreOffice did not produce a recalculated workbook')
            file.write_bytes(recalculated.read_bytes())
    values = load_workbook(file, data_only=True, read_only=True)
    errors = [f'{sheet.title}!{cell.coordinate}: {cell.value}' for sheet in values for row in sheet for cell in row if cell.data_type == 'e']
    values.close()
    if errors:
        raise RuntimeError('Formula errors: ' + '; '.join(errors[:30]))